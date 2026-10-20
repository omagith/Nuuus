"""Exports the two lecture output files, each as .docx and .pdf:

- the verbatim transcript (``[MM:SS] Лектор: текст`` lines produced by
  ``transcription.format_transcript_text``), and
- the lecture notes, written by ``structuring.structure_lecture`` in a
  small markdown subset (headings, bullets, numbered items, paragraphs,
  inline bold/italic; quotes and legacy "Уточнение по ходу лекции"
  callouts are still parsed for older notes.md files).

The markdown is parsed into blocks and inline runs here instead of being
dumped as text, so neither document ever shows raw "- " or "**" markers.
Photos of the lecture are ``![Фото 1](photos/001.jpg)`` lines (a path
relative to notes.md): each becomes the image, scaled to the page width
(and to at most ~60 % of the page height, so a portrait phone photo leaves
room for its caption), with a short caption under it.

The notes imitate a student's neat handwritten notes, so they look plain
on purpose: near-black text everywhere (headings are just bold), one
readable font, no accent colour, no shaded boxes, borders or tables. The
transcript keeps its gray timecodes and speaker labels.

PDFs are rendered from HTML. WeasyPrint gives the best typography but
needs GTK/Pango DLLs, so it stays optional and is imported lazily; when it
is missing, Qt's own QTextDocument + QPdfWriter (shipped with PyQt6, so
always present in the packaged app) renders the same content. Every piece
of text that reaches HTML goes through jinja2 autoescaping.
"""
import functools
import logging
import os
import re
import sys
import threading
from contextlib import contextmanager
from dataclasses import dataclass, field
from pathlib import Path
from typing import Callable, Iterator, List, Optional, Tuple

import jinja2
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Mm, Pt, RGBColor
from markupsafe import Markup

logger = logging.getLogger(__name__)

CALLOUT_LABEL = "Уточнение по ходу лекции"
LECTURER_LABEL = "Лектор"
APP_AUTHOR = "Нууус"  # DOCX author / PDF creator

FONT_FAMILIES = ("Calibri", "Segoe UI", "DejaVu Sans", "Arial")
CSS_FONT_STACK = '"Calibri", "Segoe UI", "DejaVu Sans", Arial, sans-serif'
DOCX_FONT = "Calibri"
BODY_PT = 11.0
PAGE_MARGIN_MM = 20.0
PDF_DPI = 300

_TEXT = RGBColor(0x1B, 0x1B, 0x1B)
_QUOTE = RGBColor(0x33, 0x33, 0x33)
_MUTED = RGBColor(0x80, 0x80, 0x80)
_AUDIENCE = RGBColor(0x5C, 0x5C, 0x5C)
# (style, size pt, space before pt, space after pt)
_HEADING_STYLES = (("Heading 1", 15, 14, 4), ("Heading 2", 13, 12, 3), ("Heading 3", 11.5, 8, 2))


class ExportError(RuntimeError):
    """Export failed; the message is user-facing (Russian)."""


@dataclass
class Block:
    kind: str  # "h1" | "h2" | "h3" | "p" | "ul" | "ol" | "quote" | "callout" | "image"
    text: str = ""  # for "image": the caption; items[0] is the image path
    items: List[str] = field(default_factory=list)


# --------------------------------------------------------------------------
# Markdown subset parsing
# --------------------------------------------------------------------------

_HEADING_RE = re.compile(r"^\s{0,3}(#{1,6})\s+(.*?)(?:\s+#+)?\s*$")
_HR_RE = re.compile(r"^\s{0,3}([-*_])(?:\s*\1){2,}\s*$")
_FENCE_RE = re.compile(r"^\s{0,3}(?:```|~~~)")
_QUOTE_RE = re.compile(r"^\s{0,3}>\s?(.*)$")
_BULLET_RE = re.compile(r"^(\s*)[-*+•]\s+(.*)$")
_NUMBERED_RE = re.compile(r"^(\s*)\d{1,4}[.)]\s+(.*)$")
_IMAGE_RE = re.compile(r'^\s{0,3}!\[(?P<alt>[^\]]*)\]\(\s*<?(?P<src>[^)>]+?)>?(?:\s+"[^"]*")?\s*\)\s*$')
# The label may be wrapped in emphasis and followed by ":" / "—" etc.
_CALLOUT_RE = re.compile(
    r"^\s*(?P<m>[*_]{0,3})\s*" + CALLOUT_LABEL + r"\s*(?P=m)?\s*[:.—–-]?\s*(?P=m)?\s*",
    re.IGNORECASE,
)
# Characters XML 1.0 forbids; python-docx raises on them and Qt/WeasyPrint
# render them as boxes. ASR/LLM output occasionally contains them.
_XML_BAD_RE = re.compile("[\x00-\x08\x0b\x0c\x0e-\x1f" + chr(0xFFFE) + chr(0xFFFF) + "]")


def _xml_safe(text: str) -> str:
    return _XML_BAD_RE.sub("", text)


def is_callout(text: str) -> bool:
    return bool(_CALLOUT_RE.match(text))


def split_callout(text: str) -> Tuple[str, str]:
    """("Уточнение по ходу лекции", rest) for callout text, ("", text) otherwise."""
    m = _CALLOUT_RE.match(text)
    if not m:
        return "", text
    return CALLOUT_LABEL, text[m.end():].strip()


class _MarkdownParser:
    """Line-based parser for the subset described in the module docstring.

    Deliberately forgiving: anything it doesn't recognize becomes a plain
    paragraph, so messy input degrades to readable text, never to an error.
    """

    def __init__(self) -> None:
        self.blocks: List[Block] = []
        self._para: List[str] = []
        self._list: Optional[Block] = None
        self._quote: List[str] = []

    def parse(self, md: str) -> List[Block]:
        for raw in (md or "").replace("\r\n", "\n").replace("\r", "\n").split("\n"):
            self._line(raw.rstrip())
        self._flush_all()
        return self.blocks

    def _line(self, line: str) -> None:
        if not line.strip():
            # Blank lines end paragraphs and quotes; lists stay open so that
            # "loose" lists (items separated by blank lines) remain one list.
            self._flush_para()
            self._flush_quote()
            return
        if _FENCE_RE.match(line):
            return  # code fences aren't part of the subset; keep their content as text
        if _HR_RE.match(line):
            self._flush_all()
            return
        m = _IMAGE_RE.match(line)
        if m:
            self._flush_all()
            self.blocks.append(Block("image", m.group("alt").strip(), [m.group("src").strip()]))
            return
        m = _HEADING_RE.match(line)
        if m and m.group(2).strip():
            self._flush_all()
            level = min(len(m.group(1)), 3)
            self.blocks.append(Block(f"h{level}", m.group(2).strip()))
            return
        m = _QUOTE_RE.match(line)
        if m:
            self._flush_para()
            self._flush_list()
            inner = m.group(1).strip()
            if not inner:
                self._flush_quote()
            else:
                if self._quote and is_callout(inner):
                    self._flush_quote()  # every callout is its own block
                self._quote.append(inner)
            return
        m = _BULLET_RE.match(line) or _NUMBERED_RE.match(line)
        if m:
            kind = "ul" if _BULLET_RE.match(line) else "ol"
            self._flush_para()
            self._flush_quote()
            if self._list is None or self._list.kind != kind:
                self._flush_list()
                self._list = Block(kind)
            self._list.items.append(m.group(2).strip())
            return
        # Plain text: indented continuation of a list item, lazy continuation
        # of a quote, or (part of) a paragraph.
        stripped = line.strip()
        if self._list is not None and self._list.items and not self._para and line.startswith("  "):
            self._list.items[-1] = f"{self._list.items[-1]} {stripped}"
            return
        if self._quote:
            self._quote.append(stripped)
            return
        self._flush_list()
        self._para.append(stripped)

    def _flush_para(self) -> None:
        if self._para:
            text = " ".join(self._para).strip()
            self._para = []
            if text:
                self.blocks.append(Block("callout" if is_callout(text) else "p", text))

    def _flush_list(self) -> None:
        if self._list is not None:
            if self._list.items:
                self.blocks.append(self._list)
            self._list = None

    def _flush_quote(self) -> None:
        if self._quote:
            text = "\n".join(self._quote).strip()
            self._quote = []
            if text:
                self.blocks.append(Block("callout" if is_callout(text) else "quote", text))

    def _flush_all(self) -> None:
        self._flush_para()
        self._flush_list()
        self._flush_quote()


def parse_markdown(md: str) -> List[Block]:
    """Parse the notes markdown subset into blocks (see module docstring)."""
    return _MarkdownParser().parse(md)


def inline_spans(text: str) -> List[Tuple[str, bool, bool]]:
    """Split ``**bold**`` / ``*italic*`` (nestable, ``***both***``) into
    (text, bold, italic) runs.

    A trimmed-down CommonMark delimiter algorithm: a run of ``*`` can open
    emphasis only if followed by non-space and close only if preceded by
    non-space, so "2 * 3" stays literal. Unmatched markers are kept as
    literal text; ``\\*`` escapes a star.
    """
    if not text:
        return []
    # Tokens: ["t", text] or ["d", remaining_count, orig_count, can_open, can_close]
    tokens: list = []
    buf: List[str] = []
    i, n = 0, len(text)
    while i < n:
        ch = text[i]
        if ch == "\\" and i + 1 < n and text[i + 1] in "*\\":
            buf.append(text[i + 1])
            i += 2
            continue
        if ch == "*":
            j = i
            while j < n and text[j] == "*":
                j += 1
            if buf:
                tokens.append(["t", "".join(buf)])
                buf = []
            before = text[i - 1] if i > 0 else " "
            after = text[j] if j < n else " "
            run = j - i
            tokens.append(["d", run, run, not after.isspace(), not before.isspace()])
            i = j
            continue
        buf.append(ch)
        i += 1
    if buf:
        tokens.append(["t", "".join(buf)])

    bold = [0] * len(tokens)
    italic = [0] * len(tokens)
    openers: List[int] = []
    for idx, tok in enumerate(tokens):
        if tok[0] != "d":
            continue
        if tok[4]:  # can close
            while tok[1] > 0:
                k = next((o for o in reversed(openers) if tokens[o][1] > 0), None)
                if k is None:
                    break
                opener = tokens[k]
                use = 2 if opener[1] >= 2 and tok[1] >= 2 else 1
                opener[1] -= use
                tok[1] -= use
                marks = bold if use == 2 else italic
                for t in range(k + 1, idx):
                    marks[t] += 1
                # Delimiters between the matched pair can no longer match.
                openers = [o for o in openers if o <= k and tokens[o][1] > 0]
        if tok[1] > 0 and tok[3]:  # can open
            openers.append(idx)

    spans: List[Tuple[str, bool, bool]] = []
    for idx, tok in enumerate(tokens):
        piece = tok[1] if tok[0] == "t" else "*" * tok[1]
        if not piece:
            continue
        style = (bold[idx] > 0, italic[idx] > 0)
        if spans and (spans[-1][1], spans[-1][2]) == style:
            spans[-1] = (spans[-1][0] + piece, style[0], style[1])
        else:
            spans.append((piece, style[0], style[1]))
    return spans


# --------------------------------------------------------------------------
# Transcript lines
# --------------------------------------------------------------------------

_TRANSCRIPT_LINE_RE = re.compile(r"^\s*\[(\d{1,2}:\d{2}(?::\d{2})?)\]\s*([^:\[\]]{1,40}?):\s?(.*)$")


@dataclass
class _TranscriptLine:
    timecode: str
    speaker: str
    text: str
    audience: bool


def _parse_transcript(transcript_text: str) -> List[object]:
    """Non-empty lines as _TranscriptLine, or plain str when the line doesn't match."""
    out: List[object] = []
    for raw in (transcript_text or "").splitlines():
        line = _xml_safe(raw.strip())
        if not line:
            continue
        m = _TRANSCRIPT_LINE_RE.match(line)
        if m:
            speaker = m.group(2).strip()
            out.append(_TranscriptLine(m.group(1), speaker, m.group(3).strip(), speaker != LECTURER_LABEL))
        else:
            out.append(line)
    return out


# --------------------------------------------------------------------------
# DOCX
# --------------------------------------------------------------------------

def _set_style_font(style, size: Optional[float] = None, color: Optional[RGBColor] = None,
                    bold: Optional[bool] = None, italic: Optional[bool] = None) -> None:
    """Pin the font to DOCX_FONT, dropping theme fonts (they'd take precedence).

    Setting ``color.rgb`` also drops the template's ``w:themeColor`` (accent
    blue), which Word would otherwise prefer over the explicit value.
    """
    rpr = style.element.get_or_add_rPr()
    rfonts = rpr.get_or_add_rFonts()
    for attr in ("w:asciiTheme", "w:hAnsiTheme", "w:eastAsiaTheme", "w:cstheme"):
        rfonts.attrib.pop(qn(attr), None)
    for attr in ("w:ascii", "w:hAnsi", "w:cs"):
        rfonts.set(qn(attr), DOCX_FONT)
    font = style.font
    if size is not None:
        font.size = Pt(size)
    if color is not None:
        font.color.rgb = color
    if bold is not None:
        font.bold = bold
    if italic is not None:
        font.italic = italic


def _drop_paragraph_decor(style) -> None:
    """Remove borders/shading a template style carries (the default Title has
    an accent-blue bottom rule)."""
    ppr = style.element.pPr
    if ppr is None:
        return
    for tag in ("w:pBdr", "w:shd"):
        old = ppr.find(qn(tag))
        if old is not None:
            ppr.remove(old)


def _new_document(title: str):
    doc = Document()
    section = doc.sections[0]
    section.page_width, section.page_height = Mm(210), Mm(297)
    for side in ("left_margin", "right_margin", "top_margin", "bottom_margin"):
        setattr(section, side, Mm(PAGE_MARGIN_MM))

    styles = doc.styles
    _set_style_font(styles["Normal"], size=BODY_PT, color=_TEXT)
    styles["Normal"].paragraph_format.space_after = Pt(4)
    styles["Normal"].paragraph_format.line_spacing = 1.2
    _set_style_font(styles["Title"], size=17, color=_TEXT, bold=True)
    _drop_paragraph_decor(styles["Title"])
    styles["Title"].paragraph_format.space_after = Pt(10)
    for name, size, before, after in _HEADING_STYLES:
        _set_style_font(styles[name], size=size, color=_TEXT, bold=True, italic=False)
        _drop_paragraph_decor(styles[name])
        styles[name].paragraph_format.space_before = Pt(before)
        styles[name].paragraph_format.space_after = Pt(after)
    for name in ("List Bullet", "List Number"):
        styles[name].paragraph_format.space_after = Pt(2)
    _set_style_font(styles["Quote"], color=_QUOTE, italic=True)
    styles["Quote"].paragraph_format.left_indent = Cm(0.6)
    styles["Quote"].paragraph_format.space_after = Pt(4)

    doc.core_properties.title = _xml_safe(title)
    doc.core_properties.author = APP_AUTHOR
    doc.core_properties.last_modified_by = APP_AUTHOR
    _add_page_number_footer(section)
    doc.add_heading(_xml_safe(title), level=0)
    return doc


def _add_page_number_footer(section) -> None:
    para = section.footer.paragraphs[0] if section.footer.paragraphs else section.footer.add_paragraph()
    para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    fld = OxmlElement("w:fldSimple")
    fld.set(qn("w:instr"), "PAGE")
    run = OxmlElement("w:r")
    rpr = OxmlElement("w:rPr")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), "808080")
    size = OxmlElement("w:sz")
    size.set(qn("w:val"), "18")
    rpr.append(color)
    rpr.append(size)
    run.append(rpr)
    t = OxmlElement("w:t")
    t.text = "1"
    run.append(t)
    fld.append(run)
    para._p.append(fld)


def _add_runs(paragraph, text: str, bold: bool = False, italic: bool = False) -> None:
    for piece, b, i in inline_spans(_xml_safe(text)):
        run = paragraph.add_run(piece)
        if b or bold:
            run.bold = True
        if i or italic:
            run.italic = True


def _add_multiline_runs(paragraph, text: str, **style) -> None:
    for n, line in enumerate(text.split("\n")):
        if n:
            paragraph.add_run().add_break()
        _add_runs(paragraph, line, **style)


def _new_numbering(doc) -> Optional[int]:
    """A fresh w:num for "List Number" that restarts at 1.

    The template's style shares one numbering instance, so without this the
    second numbered list in a document would continue from the first.
    """
    try:
        num_pr = doc.styles["List Number"].element.pPr.numPr
        numbering = doc.part.numbering_part.element
        base = numbering.num_having_numId(num_pr.numId.val)
        num = numbering.add_num(base.abstractNumId.val)
        num.add_lvlOverride(ilvl=0).add_startOverride(1)
        return num.numId
    except Exception:  # noqa: BLE001 — cosmetic; numbering simply continues
        logger.debug("Could not create a restarting numbering instance", exc_info=True)
        return None


def _set_numbering(paragraph, num_id: int) -> None:
    num_pr = paragraph._p.get_or_add_pPr().get_or_add_numPr()
    num_pr.get_or_add_ilvl().val = 0
    num_pr.get_or_add_numId().val = num_id


# --------------------------------------------------------------------------
# Images
# --------------------------------------------------------------------------

CONTENT_WIDTH_MM = 210.0 - 2 * PAGE_MARGIN_MM   # A4
IMAGE_MAX_HEIGHT_MM = 150.0                     # ~60 % of the A4 content height


def resolve_image(src: str, base_dir: Optional[Path]) -> Optional[Path]:
    """The image file a notes line points to (relative to ``base_dir``), or None if missing."""
    src = (src or "").strip()
    if not src or re.match(r"^[a-z][a-z0-9+.-]*://", src, re.IGNORECASE) and not src.lower().startswith("file:"):
        return None  # remote images are never fetched
    if src.lower().startswith("file:"):
        from urllib.parse import unquote, urlparse  # noqa: PLC0415

        src = unquote(urlparse(src).path).lstrip("/") if sys.platform == "win32" else unquote(urlparse(src).path)
    path = Path(src)
    if not path.is_absolute() and base_dir is not None:
        path = Path(base_dir) / path
    return path if path.is_file() else None


def image_display_mm(path: Path) -> Optional[Tuple[float, float]]:
    """(width, height) in mm the image is shown at: the page width at most,
    IMAGE_MAX_HEIGHT_MM high at most, never enlarged beyond ~150 dpi."""
    try:
        from PIL import Image  # noqa: PLC0415

        with Image.open(path) as im:
            w_px, h_px = im.size
    except Exception:
        logger.warning("Could not read the image %s", path, exc_info=True)
        return None
    if w_px <= 0 or h_px <= 0:
        return None
    natural_w = w_px / 150.0 * 25.4
    scale = min(1.0, CONTENT_WIDTH_MM / natural_w, IMAGE_MAX_HEIGHT_MM / (natural_w * h_px / w_px))
    width = natural_w * scale
    return width, width * h_px / w_px


def _missing_image_text(block: Block) -> str:
    return f"[{block.text or 'Фото'}: файл не найден]"


def _add_image_docx(doc, block: Block, base_dir: Optional[Path]) -> None:
    path = resolve_image(block.items[0] if block.items else "", base_dir)
    size = image_display_mm(path) if path is not None else None
    if path is None or size is None:
        para = doc.add_paragraph()
        run = para.add_run(_missing_image_text(block))
        run.italic = True
        run.font.color.rgb = _MUTED
        return
    para = doc.add_paragraph()
    para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    para.paragraph_format.space_before = Pt(6)
    para.paragraph_format.space_after = Pt(0)
    para.paragraph_format.keep_with_next = True
    try:
        para.add_run().add_picture(str(path), width=Mm(size[0]))
    except Exception:
        logger.warning("Could not embed %s into the DOCX", path, exc_info=True)
        para.add_run(_missing_image_text(block)).italic = True
        return
    if block.text:
        cap = doc.add_paragraph()
        cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
        cap.paragraph_format.space_after = Pt(8)
        run = cap.add_run(block.text)
        run.italic = True
        run.font.size = Pt(9)
        run.font.color.rgb = _MUTED


def _add_block_docx(doc, block: Block, base_dir: Optional[Path] = None) -> None:
    if block.kind == "image":
        _add_image_docx(doc, block, base_dir)
    elif block.kind in ("h1", "h2", "h3"):
        heading = doc.add_heading(level=int(block.kind[1]))
        _add_runs(heading, block.text)
    elif block.kind == "ul":
        for item in block.items:
            _add_runs(doc.add_paragraph(style="List Bullet"), item)
    elif block.kind == "ol":
        num_id = _new_numbering(doc)
        for item in block.items:
            para = doc.add_paragraph(style="List Number")
            if num_id is not None:
                _set_numbering(para, num_id)
            _add_runs(para, item)
    elif block.kind == "callout":
        # Legacy notes: an ordinary paragraph with a bold label, no box.
        label, rest = split_callout(block.text)
        para = doc.add_paragraph()
        if label:
            para.add_run(label + ": ").bold = True
        _add_multiline_runs(para, rest)
    elif block.kind == "quote":
        _add_multiline_runs(doc.add_paragraph(style="Quote"), block.text)
    else:
        _add_runs(doc.add_paragraph(), block.text)


@contextmanager
def _atomic_target(out_path: Path) -> Iterator[Path]:
    """Write to a temp file next to the target and swap it in at the end.

    A half-written file never replaces a good one, and a target that is open
    in Word/Acrobat (locked on Windows) produces a clear message.
    """
    out_path = Path(out_path)
    try:
        out_path.parent.mkdir(parents=True, exist_ok=True)
    except OSError as exc:
        raise ExportError(f"Не удалось создать папку {out_path.parent}: {exc}") from exc
    tmp = out_path.with_name(f"{out_path.stem}.tmp{out_path.suffix}")
    try:
        yield tmp
        os.replace(tmp, out_path)
    except PermissionError as exc:
        raise ExportError(
            f"Не удалось сохранить «{out_path.name}»: файл открыт в другой программе или нет доступа. "
            "Закройте его и повторите."
        ) from exc
    except OSError as exc:
        raise ExportError(f"Не удалось сохранить «{out_path.name}»: {exc}") from exc
    finally:
        try:
            if tmp.exists():
                tmp.unlink()
        except OSError:
            pass


def _save_docx(doc, out_path: Path) -> None:
    with _atomic_target(out_path) as tmp:
        doc.save(str(tmp))
    logger.info("Saved %s", out_path)


def export_transcript_docx(transcript_text: str, out_path: Path, title: str) -> None:
    doc = _new_document(title)
    lines = _parse_transcript(transcript_text)
    if not lines:
        doc.add_paragraph("Речь в записи не распознана.")
    for line in lines:
        para = doc.add_paragraph()
        para.paragraph_format.space_after = Pt(3)
        if isinstance(line, str):
            para.add_run(line)
            continue
        tc = para.add_run(f"{line.timecode}  ")
        tc.font.size = Pt(8.5)
        tc.font.color.rgb = _MUTED
        speaker = para.add_run(f"{line.speaker}: ")
        speaker.bold = True
        text = para.add_run(line.text)
        if line.audience:
            para.paragraph_format.left_indent = Cm(0.8)
            speaker.font.color.rgb = _AUDIENCE
            text.italic = True
            text.font.color.rgb = _AUDIENCE
    _save_docx(doc, out_path)


def export_notes_docx(markdown: str, out_path: Path, title: str, base_dir: Optional[Path] = None) -> None:
    """``base_dir``: where relative image paths point from (default: the output folder)."""
    doc = _new_document(title)
    blocks = parse_markdown(markdown)
    if not blocks:
        doc.add_paragraph("Конспект пуст.")
    base = Path(base_dir) if base_dir is not None else Path(out_path).parent
    for block in blocks:
        _add_block_docx(doc, block, base)
    _save_docx(doc, out_path)


# --------------------------------------------------------------------------
# HTML
# --------------------------------------------------------------------------

# Lengths go through L() (vertical) / H() (horizontal): points for WeasyPrint.
# QTextDocument ignores pt in margins; it scales vertical px from 96 dpi to the
# device but uses horizontal margins/paddings as raw device pixels.
_BASE_TEMPLATE = """<!DOCTYPE html>
<html lang="ru"><head><meta charset="utf-8"><title>{{ title }}</title>
<style>
{% if not qt %}@page { size: A4; margin: {{ margin_mm }}mm; }
{% endif %}body { font-family: {{ font_stack }}; font-size: {{ body_pt }}pt; color: #1b1b1b;{% if not qt %} line-height: 1.4;{% endif %} }
h1.title { font-size: 17pt; color: #1b1b1b; margin-top: 0; margin-bottom: {{ L(12) }}; }
h1 { font-size: 15pt; color: #1b1b1b; margin-top: {{ L(14) }}; margin-bottom: {{ L(5) }}; }
h2 { font-size: 13pt; color: #1b1b1b; margin-top: {{ L(13) }}; margin-bottom: {{ L(4) }}; }
h3 { font-size: 11.5pt; color: #1b1b1b; margin-top: {{ L(9) }}; margin-bottom: {{ L(3) }}; }
p { margin-top: {{ L(3) }}; margin-bottom: {{ L(5) }}; }
ul, ol { margin-top: {{ L(2) }}; margin-bottom: {{ L(6) }}; }
li { margin-top: {{ L(2) }}; margin-bottom: {{ L(2) }}; }
.quote { color: #333333; font-style: italic; margin-left: {{ H(14) }}; }
p.line { margin-top: {{ L(1.5) }}; margin-bottom: {{ L(1.5) }}; }
.tc { color: #808080; font-size: 8.5pt; }
.spk { font-weight: bold; }
p.aud { color: #5c5c5c; margin-left: {{ H(16) }}; }
.aud .txt { font-style: italic; }
p.img { text-align: center; margin-top: {{ L(8) }}; margin-bottom: {{ L(2) }};{% if not qt %} page-break-inside: avoid; page-break-after: avoid;{% endif %} }
{% if not qt %}p.img img { max-width: 100%; }
{% endif %}p.cap { text-align: center; color: #808080; font-size: 9pt; font-style: italic; margin-top: {{ L(2) }}; margin-bottom: {{ L(10) }}; }
p.missing { color: #808080; font-style: italic; }
</style></head><body>
<h1 class="title">{{ title }}</h1>
{% block content %}{% endblock %}
</body></html>
"""

_MACROS = """{% macro spans(items) -%}
{%- for text, bold, italic in items -%}
{%- if bold %}<strong>{% endif %}{% if italic %}<em>{% endif %}{{ text }}{% if italic %}</em>{% endif %}{% if bold %}</strong>{% endif -%}
{%- endfor -%}
{%- endmacro %}
{% macro lines(items) -%}
{%- for line in items %}{% if not loop.first %}<br>{% endif %}{{ spans(line) }}{% endfor -%}
{%- endmacro %}
"""

_NOTES_TEMPLATE = """{% extends "base" %}{% import "macros" as m %}
{% block content %}
{% for b in blocks %}
{% if b.kind in ("h1", "h2", "h3") %}<{{ b.kind }}>{{ m.spans(b.spans) }}</{{ b.kind }}>
{% elif b.kind == "ul" %}<ul>{% for item in b.entries %}<li>{{ m.spans(item) }}</li>{% endfor %}</ul>
{% elif b.kind == "ol" %}<ol>{% for item in b.entries %}<li>{{ m.spans(item) }}</li>{% endfor %}</ol>
{% elif b.kind == "callout" %}<p>{% if b.label %}<strong>{{ b.label }}:</strong> {% endif %}{{ m.lines(b.lines) }}</p>
{% elif b.kind == "quote" %}<p class="quote">{{ m.lines(b.lines) }}</p>
{% elif b.kind == "image" %}{% if b.src %}<p class="img" align="center"><img src="{{ b.src }}" width="{{ b.width }}" height="{{ b.height }}"></p>
{% if b.caption %}<p class="cap" align="center">{{ b.caption }}</p>
{% endif %}{% else %}<p class="missing">{{ b.caption }}</p>
{% endif %}
{% else %}<p>{{ m.lines(b.lines) }}</p>
{% endif %}
{% endfor %}
{% endblock %}
"""

_TRANSCRIPT_TEMPLATE = """{% extends "base" %}
{% block content %}
{% for line in lines %}
{% if line.speaker is defined %}<p class="line{% if line.audience %} aud{% endif %}"><span class="tc">{{ line.timecode }}</span>&nbsp;&nbsp;<span class="spk">{{ line.speaker }}:</span> <span class="txt">{{ line.text }}</span></p>
{% else %}<p class="line">{{ line }}</p>
{% endif %}
{% else %}<p>Речь в записи не распознана.</p>
{% endfor %}
{% endblock %}
"""

_JINJA = jinja2.Environment(
    loader=jinja2.DictLoader({
        "base": _BASE_TEMPLATE,
        "macros": _MACROS,
        "notes": _NOTES_TEMPLATE,
        "transcript": _TRANSCRIPT_TEMPLATE,
    }),
    autoescape=True,
    trim_blocks=True,
    lstrip_blocks=True,
)


def _template_context(title: str, qt: bool) -> dict:
    return {
        "title": _xml_safe(title),
        "qt": qt,
        "margin_mm": PAGE_MARGIN_MM,
        # <style> content is raw text (entities are not decoded there), so the
        # quoted font names must not be escaped.
        "font_stack": Markup(CSS_FONT_STACK),
        "body_pt": f"{BODY_PT:g}",
        "L": (lambda pt: f"{pt * 96 / 72.0:.0f}px") if qt else (lambda pt: f"{pt:g}pt"),
        "H": (lambda pt: f"{pt * PDF_DPI / 72.0:.0f}px") if qt else (lambda pt: f"{pt:g}pt"),
    }


def _block_view(block: Block, base_dir: Optional[Path] = None) -> dict:
    text = _xml_safe(block.text)
    view = {"kind": block.kind}
    if block.kind == "image":
        path = resolve_image(block.items[0] if block.items else "", base_dir)
        size = image_display_mm(path) if path is not None else None
        if path is None or size is None:
            view.update(src="", caption=_missing_image_text(block))
        else:
            # Sizes in CSS px (96 dpi): WeasyPrint uses them as is, Qt scales them to the device.
            view.update(src=path.resolve().as_uri(), caption=text,
                        width=max(1, round(size[0] / 25.4 * 96)), height=max(1, round(size[1] / 25.4 * 96)))
    elif block.kind in ("h1", "h2", "h3"):
        view["spans"] = inline_spans(text)
    elif block.kind in ("ul", "ol"):
        view["entries"] = [inline_spans(_xml_safe(item)) for item in block.items]
    elif block.kind == "callout":
        label, rest = split_callout(text)
        view["label"] = label
        view["lines"] = [inline_spans(line) for line in rest.split("\n")]
    else:
        view["lines"] = [inline_spans(line) for line in text.split("\n")]
    return view


def render_notes_html(markdown: str, title: str, qt: bool = False, base_dir: Optional[Path] = None) -> str:
    """Notes as a standalone HTML page (``qt=True``: markup tuned for QTextDocument).
    Images are referenced by absolute file:// URLs (``base_dir`` resolves relative ones)."""
    blocks = [_block_view(b, base_dir) for b in parse_markdown(markdown)]
    return _JINJA.get_template("notes").render(blocks=blocks, **_template_context(title, qt))


def render_transcript_html(transcript_text: str, title: str, qt: bool = False) -> str:
    lines = _parse_transcript(transcript_text)
    return _JINJA.get_template("transcript").render(lines=lines, **_template_context(title, qt))


# --------------------------------------------------------------------------
# PDF
# --------------------------------------------------------------------------

def _import_weasyprint():
    """WeasyPrint module or None. Import errors include OSError when the
    GTK/Pango DLLs are missing, so anything raised means "unavailable"."""
    try:
        import weasyprint  # noqa: PLC0415 — optional, heavy, needs GTK
        return weasyprint
    except Exception as exc:  # noqa: BLE001
        logger.info("WeasyPrint unavailable (%s: %s)", type(exc).__name__, exc)
        return None


def _qt_pdf_importable() -> bool:
    try:
        from PyQt6.QtGui import QPdfWriter, QTextDocument  # noqa: F401, PLC0415
        return True
    except Exception as exc:  # noqa: BLE001
        logger.info("Qt PDF backend unavailable (%s: %s)", type(exc).__name__, exc)
        return False


@functools.lru_cache(maxsize=1)
def pdf_backend() -> str:
    """Best available PDF engine: "weasyprint", "qt" or "none" (probed once)."""
    if _import_weasyprint() is not None:
        return "weasyprint"
    if _qt_pdf_importable():
        return "qt"
    return "none"


def _pdf_weasyprint(html: str, out_path: Path) -> None:
    weasyprint = _import_weasyprint()
    if weasyprint is None:
        raise RuntimeError("WeasyPrint недоступен")
    with _atomic_target(out_path) as tmp:
        weasyprint.HTML(string=html).write_pdf(str(tmp))


_qt_app = None  # keeps a lazily created QGuiApplication alive


def _ensure_qt_gui_app() -> None:
    """QTextDocument needs a QGuiApplication (fonts). Reuse the app's one; create
    one only on the main thread (Qt forbids creating it elsewhere)."""
    global _qt_app
    from PyQt6.QtCore import QCoreApplication  # noqa: PLC0415
    from PyQt6.QtGui import QGuiApplication  # noqa: PLC0415

    app = QCoreApplication.instance()
    if app is None:
        if threading.current_thread() is not threading.main_thread():
            raise ExportError(
                "Не удалось создать PDF: движку Qt нужно запущенное приложение "
                "(экспорт вызван из фонового потока без QApplication)."
            )
        _qt_app = QGuiApplication(sys.argv[:1] or [APP_AUTHOR])
    elif not isinstance(app, QGuiApplication):
        raise ExportError(
            "Не удалось создать PDF: движку Qt нужно графическое приложение (QGuiApplication)."
        )


_qt_fonts_lock = threading.Lock()
_qt_fonts_checked = False


def _ensure_qt_fonts() -> None:
    """Make sure QTextDocument has a Cyrillic-capable font.

    The normal "windows" platform plugin sees all system fonts, but Qt's
    headless "offscreen" plugin (tests, services) starts with an empty font
    database and would render every glyph as a box; register the fonts from
    the Windows font folder in that case.
    """
    global _qt_fonts_checked
    with _qt_fonts_lock:
        if _qt_fonts_checked:
            return
        _qt_fonts_checked = True
        from PyQt6.QtGui import QFontDatabase  # noqa: PLC0415

        if any(f in QFontDatabase.families() for f in FONT_FAMILIES):
            return
        fonts_dir = Path(os.environ.get("WINDIR", r"C:\Windows")) / "Fonts"
        loaded = 0
        for name in ("calibri.ttf", "calibrib.ttf", "calibrii.ttf", "calibriz.ttf",
                     "segoeui.ttf", "segoeuib.ttf", "segoeuii.ttf", "segoeuiz.ttf",
                     "arial.ttf", "arialbd.ttf", "ariali.ttf", "arialbi.ttf"):
            path = fonts_dir / name
            if path.exists() and QFontDatabase.addApplicationFont(str(path)) >= 0:
                loaded += 1
        logger.info("Qt font database was empty; registered %d font files from %s", loaded, fonts_dir)


def _qt_print(html: str, target: Path, title: str) -> None:
    """Lay out HTML with QTextDocument and paginate it onto QPdfWriter by hand,
    which gives exact A4 margins and "N / M" page numbers in the footer.

    Kept separate so every Qt object is released when it returns or raises
    (QPdfWriter holds the file open until it is destroyed).
    """
    from PyQt6.QtCore import QMarginsF, QRectF, QSizeF, Qt  # noqa: PLC0415
    from PyQt6.QtGui import (  # noqa: PLC0415
        QColor, QFont, QPageLayout, QPageSize, QPainter, QPdfWriter, QTextDocument,
    )

    writer = QPdfWriter(str(target))
    doc = QTextDocument()
    painter = QPainter()
    try:
        writer.setPageSize(QPageSize(QPageSize.PageSizeId.A4))
        writer.setPageMargins(QMarginsF(0, 0, 0, 0), QPageLayout.Unit.Millimeter)
        writer.setResolution(PDF_DPI)
        writer.setTitle(title)
        writer.setCreator(APP_AUTHOR)

        page_w, page_h = float(writer.width()), float(writer.height())
        margin = PAGE_MARGIN_MM / 25.4 * PDF_DPI
        content_w, content_h = page_w - 2 * margin, page_h - 2 * margin

        doc.documentLayout().setPaintDevice(writer)  # pt sizes -> 300 dpi device pixels
        font = QFont()
        font.setFamilies(list(FONT_FAMILIES))
        font.setPointSizeF(BODY_PT)
        doc.setDefaultFont(font)
        doc.setDocumentMargin(0)
        doc.setIndentWidth(18 * 96 / 72.0)  # CSS px; Qt scales them to the device
        doc.setHtml(html)
        doc.setPageSize(QSizeF(content_w, content_h))
        pages = max(1, doc.pageCount())

        if not painter.begin(writer):
            raise ExportError(f"Не удалось открыть «{target.name}» для записи PDF.")
        footer_font = QFont(font)
        footer_font.setPointSizeF(8)
        for page in range(pages):
            if page:
                writer.newPage()
            painter.save()
            painter.translate(margin, margin - page * content_h)
            # One device pixel short: a photo moved to the next page starts
            # exactly at the boundary and would leave its top row here.
            doc.drawContents(painter, QRectF(0, page * content_h, content_w, content_h - 1))
            painter.restore()
            painter.setFont(footer_font)
            painter.setPen(QColor("#808080"))
            painter.drawText(
                QRectF(0, page_h - margin * 0.75, page_w, margin * 0.5),
                Qt.AlignmentFlag.AlignCenter,
                f"{page + 1} / {pages}",
            )
    finally:
        if painter.isActive():
            painter.end()  # finalizes the PDF
        del painter, doc, writer


def _pdf_qt(html: str, out_path: Path, title: str) -> None:
    _ensure_qt_gui_app()
    _ensure_qt_fonts()
    with _atomic_target(out_path) as tmp:
        _qt_print(html, tmp, title)


def _write_pdf(render: Callable[[bool], str], out_path: Path, title: str) -> None:
    backend = pdf_backend()
    problems: List[str] = []
    if backend == "weasyprint":
        try:
            _pdf_weasyprint(render(False), out_path)
            logger.info("Saved %s (WeasyPrint)", out_path)
            return
        except ExportError:
            raise  # file locked etc. — Qt would hit the same problem
        except Exception as exc:  # noqa: BLE001 — fall back to Qt below
            logger.warning("WeasyPrint failed for %s, falling back to Qt", out_path, exc_info=True)
            problems.append(f"WeasyPrint: {exc}")
        backend = "qt" if _qt_pdf_importable() else "none"
    if backend == "qt":
        try:
            _pdf_qt(render(True), out_path, title)
            logger.info("Saved %s (Qt)", out_path)
            return
        except ExportError:
            raise
        except Exception as exc:  # noqa: BLE001
            logger.exception("Qt PDF export failed for %s", out_path)
            problems.append(f"Qt: {exc}")
    detail = "; ".join(problems) if problems else "нет ни WeasyPrint, ни модуля печати Qt"
    raise ExportError(f"Не удалось создать PDF «{Path(out_path).name}» ({detail}).")


def export_transcript_pdf(transcript_text: str, out_path: Path, title: str) -> None:
    _write_pdf(lambda qt: render_transcript_html(transcript_text, title, qt=qt), out_path, title)


def export_notes_pdf(markdown: str, out_path: Path, title: str, base_dir: Optional[Path] = None) -> None:
    """``base_dir``: where relative image paths point from (default: the output folder)."""
    base = Path(base_dir) if base_dir is not None else Path(out_path).parent
    _write_pdf(lambda qt: render_notes_html(markdown, title, qt=qt, base_dir=base), out_path, title)
