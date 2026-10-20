"""LecturePipeline with every stage function faked (no models, GPU or Ollama)."""
import json
import threading

import numpy as np
import pytest
import soundfile as sf

from app.core import lecture_pipeline as lp
from app.core.errors import Cancelled
from app.core.lecture_pipeline import (
    CACHE_VERSION,
    STAGE_LABELS,
    STAGES,
    LecturePipeline,
    LectureResult,
    recording_label,
)
from app.core.models_config import llm_options, llm_spec
from app.lecture import audio_io, diarization, export, preprocessing, transcription
from app.lecture.diarization import DiarizedSegment
from app.llm.ollama_client import OllamaError, OllamaModelMissingError, OllamaUnavailableError

SR = 16000
NOTES = "## Пределы [00:00–00:03]\n- **Предел** — значение, к которому стремится функция"


# --------------------------------------------------------------------------- fakes

class FakeSTT:
    checkpoint = "fake"

    def __init__(self):
        self.batches = []

    def transcribe_batch(self, clips, sample_rate=SR, batch_size=8):
        self.batches.append(len(clips))
        return [f"фраза длиной {len(c) / sample_rate:.1f} с" for c in clips]


class FakeLLM:
    def __init__(self, reply=NOTES):
        self.reply = reply
        self.ready_error = None
        self.stream_error = None
        self.calls = []
        self.ensure_calls = []

    def ensure_ready(self, model):
        self.ensure_calls.append(model)
        if self.ready_error is not None:
            raise self.ready_error

    def stream_generate(self, model, prompt, system=None, on_token=None, options=None, cancel_event=None,
                        **_kw):
        self.calls.append({"model": model, "prompt": prompt, "options": options})
        if self.stream_error is not None:
            raise self.stream_error
        if on_token:
            on_token(self.reply)
        return self.reply


class Stages:
    """Records calls to the faked stage functions."""

    def __init__(self):
        self.calls = []
        self.diarize_error = None
        self.diarize_hook = None  # called inside diarize with on_progress
        self.preprocess_error = None
        self.turns = [
            DiarizedSegment(0.0, 1.0, "SPEAKER_00"),
            DiarizedSegment(1.2, 2.0, "SPEAKER_00"),
            DiarizedSegment(2.1, 2.5, "SPEAKER_01"),
            DiarizedSegment(2.6, 3.0, "SPEAKER_00"),
        ]
        self.stt = FakeSTT()
        self.devices = {}      # stage -> device it was asked to use
        self.released = []     # (checkpoint, device) passed to release_engine
        self.planned = []      # (task, tier, device, stt_checkpoint) passed to plan_llm_options

    def preprocess(self, audio, sample_rate, do_denoise=False, do_dereverb=False, on_progress=None,
                   on_warning=None, cancel_event=None, device=None):
        self.calls.append(("preprocess", do_denoise, do_dereverb))
        self.devices["preprocess"] = device
        assert isinstance(audio, np.ndarray) and sample_rate == SR
        if self.preprocess_error is not None:
            raise self.preprocess_error
        if on_progress:
            on_progress(0.5)
            on_progress(1.0)
        if do_denoise and on_warning:
            on_warning("Шумоподавление пропущено: не установлен модуль df")
        return audio * 0.5

    def diarize(self, audio, hf_token, sample_rate=SR, device=None, on_progress=None):
        self.calls.append(("diarize", hf_token))
        self.devices["diarize"] = device
        assert isinstance(audio, np.ndarray)
        if on_progress:
            on_progress(0.4)
        if self.diarize_hook:
            self.diarize_hook(on_progress)
        if self.diarize_error is not None:
            raise self.diarize_error
        return list(self.turns)

    def vad_segments(self, audio, sample_rate=SR, speaker="SPEAKER_00", on_progress=None, **_kw):
        self.calls.append(("vad",))
        if on_progress:
            on_progress(0.5)
            on_progress(1.0)
        return [DiarizedSegment(0.0, 2.9, "SPEAKER_00")]

    def get_engine(self, checkpoint, device=None):
        self.calls.append(("get_engine", checkpoint))
        self.devices["stt"] = device
        return self.stt

    def release_engine(self, checkpoint, device=None):
        self.released.append((checkpoint, device))

    def plan_llm_options(self, task, tier, device="gpu", stt_checkpoint=None, client=None):
        self.planned.append((task, tier, device, stt_checkpoint))
        return llm_options(task, tier, device)

    def names(self):
        return [c[0] for c in self.calls]


class Events:
    def __init__(self):
        self.items = []

    def cb(self, name):
        def _cb(*args):
            self.items.append((name,) + args)
        return _cb

    def of(self, name):
        return [e[1:] for e in self.items if e[0] == name]

    def terminal(self):
        return [e for e in self.items if e[0] in ("done", "error", "cancelled")]


@pytest.fixture(autouse=True)
def _short_fake_transcripts_count_as_material(monkeypatch):
    """The fake STT says "фраза длиной 2.0 с": too few content words for
    structuring's "is there anything to take notes on" gate, which would skip
    the LLM. These tests are about the pipeline, so any content word counts."""
    from app.lecture import structuring

    monkeypatch.setattr(structuring, "MIN_CONTENT_WORDS", 1)


@pytest.fixture
def env(tmp_path, monkeypatch):
    """Audio file + faked stage functions; outputs go to tmp_path/out/<stem>."""
    stages = Stages()
    monkeypatch.setattr(preprocessing, "preprocess", stages.preprocess)
    monkeypatch.setattr(diarization, "diarize", stages.diarize)
    monkeypatch.setattr(diarization, "vad_segments", stages.vad_segments)
    monkeypatch.setattr(lp, "get_engine", stages.get_engine)
    monkeypatch.setattr(lp, "release_engine", stages.release_engine)
    monkeypatch.setattr(lp, "plan_llm_options", stages.plan_llm_options)
    out_root = tmp_path / "out"
    monkeypatch.setattr(audio_io, "output_dir_for", lambda p: out_root / p.stem)
    monkeypatch.delenv("HF_TOKEN", raising=False)

    pdf_calls = []

    def fake_pdf(kind):
        def _pdf(text, out_path, title):
            pdf_calls.append((kind, out_path.name, title))
            out_path.write_bytes(b"%PDF-1.4 fake")
        return _pdf

    monkeypatch.setattr(export, "export_transcript_pdf", fake_pdf("transcript"))
    monkeypatch.setattr(export, "export_notes_pdf", fake_pdf("notes"))
    stages.pdf_calls = pdf_calls

    audio = tmp_path / "raw_20261003_201235.wav"
    rng = np.random.default_rng(0)
    sf.write(str(audio), (rng.standard_normal(3 * SR) * 0.1).astype(np.float32), SR, subtype="PCM_16")
    stages.audio = audio
    stages.out_dir = out_root / audio.stem
    return stages


def make(env, llm=None, **kw):
    events = Events()
    llm = llm or FakeLLM()
    pipeline = LecturePipeline(
        stt_tier=kw.pop("stt_tier", "light"), llm_tier=kw.pop("llm_tier", "light"), llm_client=llm,
        on_stage=events.cb("stage"), on_progress=events.cb("progress"), on_warning=events.cb("warning"),
        on_done=events.cb("done"), on_error=events.cb("error"), on_cancelled=events.cb("cancelled"),
        **kw,
    )
    return pipeline, events, llm


def result_of(events) -> LectureResult:
    terminal = events.terminal()
    assert len(terminal) == 1, terminal
    assert terminal[0][0] == "done", terminal
    return terminal[0][1]


# --------------------------------------------------------------------------- happy path

def test_happy_path_order_files_and_cache(env):
    pipeline, events, llm = make(env)
    pipeline.run(env.audio)
    result = result_of(events)

    assert [s for (s,) in events.of("stage")] == STAGES
    # Every stage's bar starts at 0 and ends at 1, never going backwards.
    for stage in STAGES:
        values = [v for s, v in events.of("progress") if s == stage]
        assert values[0] == 0.0 and values[-1] == 1.0
        assert values == sorted(values) and len(values) == len(set(values))
    # Stage events precede their progress events, in stage order.
    order = [e[1] for e in events.items if e[0] in ("stage", "progress")]
    assert [s for i, s in enumerate(order) if i == 0 or order[i - 1] != s] == STAGES

    out = env.out_dir
    assert result.out_dir == out and result.audio_path == env.audio
    assert result.warnings == [] and events.of("warning") == []
    for key, name in audio_io.OUTPUT_FILES.items():
        assert getattr(result, key) == out / name
        assert (out / name).is_file(), name

    assert env.names() == ["preprocess", "diarize", "get_engine"]
    assert env.calls[2] == ("get_engine", "v3_e2e_ctc")
    assert (out / "clean.wav").is_file()
    clean_info = sf.info(str(out / "clean.wav"))
    assert clean_info.samplerate == SR and clean_info.subtype == "PCM_16"

    transcript = (out / "transcript.txt").read_text(encoding="utf-8")
    lines = transcript.strip().split("\n")
    # merge_segments joined the two first lecturer turns (gap 0.2 s).
    assert lines == [
        "[00:00] Лектор: фраза длиной 2.0 с",
        "[00:02] Аудитория: фраза длиной 0.4 с",
        "[00:02] Лектор: фраза длиной 0.4 с",
    ]
    notes = (out / "notes.md").read_text(encoding="utf-8")
    assert "Пределы" in notes  # the notes format itself is structuring's business

    cache = json.loads((out / "segments.json").read_text(encoding="utf-8"))
    assert cache["version"] == CACHE_VERSION
    assert cache["stt_checkpoint"] == "v3_e2e_ctc"
    assert cache["denoise"] is False and cache["dereverb"] is False
    assert cache["lecturer"] == "SPEAKER_00"
    assert cache["diarization"] == "pyannote"
    assert cache["warnings"] == []
    assert [s["speaker"] for s in cache["segments"]] == ["SPEAKER_00", "SPEAKER_01", "SPEAKER_00"]

    assert llm.ensure_calls == [llm_spec("lecture", "light").ollama_tag]
    # Planned once for the whole stage (after the STT model was released), same dict for every chunk.
    assert env.planned == [("lecture", "light", "gpu", None)]
    assert llm.calls[0]["options"] == llm_options("lecture", "light", "gpu")
    assert env.released and env.released[0] == ("v3_e2e_ctc", "gpu")
    titles = {kind: title for kind, _name, title in env.pdf_calls}
    assert titles == {"transcript": "Расшифровка лекции — 03.10.2026 20:12",
                      "notes": "Конспект лекции — 03.10.2026 20:12"}


def test_real_docx_export_and_language_switch(env):
    pipeline, events, llm = make(env, language="multi", stt_tier="power")
    pipeline.run(env.audio)
    result = result_of(events)
    assert env.calls[2] == ("get_engine", "multilingual_large_ctc")
    from docx import Document

    doc = Document(str(result.transcript_docx))
    text = "\n".join(p.text for p in doc.paragraphs)
    assert "Расшифровка лекции — 03.10.2026 20:12" in text
    assert "Лектор: фраза длиной 2.0 с" in text
    notes = "\n".join(p.text for p in Document(str(result.notes_docx)).paragraphs)
    assert "Пределы" in notes and "**" not in notes


def test_cache_hit_skips_expensive_stages(env):
    pipeline, events, llm = make(env)
    pipeline.run(env.audio)
    result_of(events)
    env.calls.clear()
    (env.out_dir / "transcript.txt").unlink()

    pipeline2, events2, llm2 = make(env)
    pipeline2.run(env.audio)
    result = result_of(events2)
    assert env.calls == []  # no preprocess / diarize / model load
    assert [s for (s,) in events2.of("stage")] == STAGES
    for stage in STAGES[:3]:
        assert [v for s, v in events2.of("progress") if s == stage] == [0.0, 1.0]
    assert result.transcript_txt.is_file()
    assert "Лектор: фраза длиной 2.0 с" in result.transcript_txt.read_text(encoding="utf-8")
    assert llm2.calls  # the notes are still (re)generated


@pytest.mark.parametrize("change", [
    {"stt_tier": "power"},
    {"denoise": True},
    {"dereverb": True},
])
def test_cache_mismatch_reprocesses(env, change):
    pipeline, events, _llm = make(env)
    pipeline.run(env.audio)
    result_of(events)
    env.calls.clear()
    pipeline2, events2, _llm2 = make(env, **change)
    pipeline2.run(env.audio)
    result_of(events2)
    assert env.names()[:2] == ["preprocess", "diarize"]


def test_corrupt_cache_is_ignored(env):
    env.out_dir.mkdir(parents=True)
    (env.out_dir / "segments.json").write_text("{not json", encoding="utf-8")
    pipeline, events, _llm = make(env)
    pipeline.run(env.audio)
    result_of(events)
    assert env.names() == ["preprocess", "diarize", "get_engine"]
    assert json.loads((env.out_dir / "segments.json").read_text(encoding="utf-8"))["version"] == CACHE_VERSION


def test_cache_with_bad_segments_is_ignored(env):
    env.out_dir.mkdir(parents=True)
    (env.out_dir / "segments.json").write_text(json.dumps({
        "version": CACHE_VERSION, "stt_checkpoint": "v3_e2e_ctc", "denoise": False, "dereverb": False,
        "lecturer": "SPEAKER_00", "diarization": "pyannote", "segments": [{"start": "x"}], "warnings": [],
    }), encoding="utf-8")
    pipeline, events, _llm = make(env)
    pipeline.run(env.audio)
    result_of(events)
    assert "diarize" in env.names()


def test_cache_warnings_are_repeated(env):
    env.diarize_error = RuntimeError("нет токена")
    pipeline, events, _llm = make(env)
    pipeline.run(env.audio)
    first = result_of(events).warnings
    pipeline2, events2, _llm2 = make(env)
    pipeline2.run(env.audio)
    assert result_of(events2).warnings == first
    assert len(first) == 1 and "Диаризация недоступна" in first[0]


def test_vad_cache_is_redone_once_a_token_is_available(env, monkeypatch):
    env.diarize_error = RuntimeError("Не удалось загрузить модель диаризации pyannote.")
    pipeline, events, _llm = make(env)
    pipeline.run(env.audio)
    result_of(events)
    env.calls.clear()
    env.diarize_error = None
    monkeypatch.setenv("HF_TOKEN", "hf_test")
    pipeline2, events2, _llm2 = make(env)
    pipeline2.run(env.audio)
    result = result_of(events2)
    assert ("diarize", "hf_test") in env.calls
    assert result.warnings == []
    cache = json.loads((env.out_dir / "segments.json").read_text(encoding="utf-8"))
    assert cache["diarization"] == "pyannote" and cache["hf_token_set"] is True


def test_changed_audio_file_invalidates_cache(env):
    pipeline, events, _llm = make(env)
    pipeline.run(env.audio)
    result_of(events)
    sf.write(str(env.audio), np.zeros(4 * SR, dtype=np.float32), SR, subtype="PCM_16")
    env.calls.clear()
    pipeline2, events2, _llm2 = make(env)
    pipeline2.run(env.audio)
    result_of(events2)
    assert "diarize" in env.names()


# --------------------------------------------------------------------------- degraded runs

def test_diarization_failure_falls_back_to_vad(env):
    env.diarize_error = RuntimeError(
        "Не удалось загрузить модель диаризации pyannote. Проверьте токен Hugging Face. "
        "Подробности: pyannote/x: 401 Client Error")
    pipeline, events, _llm = make(env)
    pipeline.run(env.audio)
    result = result_of(events)
    expected = ("Диаризация недоступна (Не удалось загрузить модель диаризации pyannote. Проверьте токен "
                "Hugging Face). Вся речь будет считаться речью лектора.")
    assert result.warnings == [expected]
    assert events.of("warning") == [(expected,)]
    assert env.names() == ["preprocess", "diarize", "vad", "get_engine"]
    values = [v for s, v in events.of("progress") if s == "diarization"]
    assert values == sorted(values) and values[-1] == 1.0
    cache = json.loads((env.out_dir / "segments.json").read_text(encoding="utf-8"))
    assert cache["diarization"] == "vad" and cache["warnings"] == [expected]
    assert (env.out_dir / "transcript.txt").read_text(encoding="utf-8").startswith("[00:00] Лектор:")


def test_preprocessing_warnings_are_forwarded_and_cached(env):
    pipeline, events, _llm = make(env, denoise=True)
    pipeline.run(env.audio)
    result = result_of(events)
    assert result.warnings == ["Шумоподавление пропущено: не установлен модуль df"]
    assert env.calls[0] == ("preprocess", True, False)
    cache = json.loads((env.out_dir / "segments.json").read_text(encoding="utf-8"))
    assert cache["denoise"] is True and cache["warnings"] == result.warnings


@pytest.mark.parametrize("error", [
    OllamaUnavailableError("Ollama не запущена или недоступна по адресу http://localhost:11434. "
                           "Запустите приложение Ollama."),
    OllamaModelMissingError("lakomoor/vikhr-llama-3.2-1b-instruct:1b"),
])
def test_ollama_not_ready_exports_transcript_only(env, error):
    llm = FakeLLM()
    llm.ready_error = error
    pipeline, events, _llm = make(env, llm=llm)
    pipeline.run(env.audio)
    result = result_of(events)
    assert llm.calls == []
    assert result.notes_docx is None and result.notes_pdf is None and result.notes_md is None
    assert result.transcript_docx.is_file() and result.transcript_txt.is_file() and result.transcript_pdf.is_file()
    assert not (env.out_dir / "notes.md").exists()
    assert len(result.warnings) == 1
    warning = result.warnings[0]
    assert warning.startswith("Конспект не составлен: " + str(error).rstrip("."))
    assert "обработайте запись ещё раз" in warning
    assert [s for (s,) in events.of("stage")] == STAGES


def test_structuring_error_skips_notes(env):
    llm = FakeLLM()
    llm.stream_error = OllamaError("Ошибка Ollama: CUDA out of memory")
    pipeline, events, _llm = make(env, llm=llm)
    pipeline.run(env.audio)
    result = result_of(events)
    assert result.notes_docx is None and result.transcript_docx is not None
    assert result.warnings[0].startswith("Конспект не составлен: Ошибка Ollama: CUDA out of memory.")


def test_no_speech_makes_empty_notes_without_llm(env):
    env.turns = []
    llm = FakeLLM()
    llm.ready_error = OllamaUnavailableError("не запущена")  # must not even be asked
    pipeline, events, _llm = make(env, llm=llm)
    pipeline.run(env.audio)
    result = result_of(events)
    assert llm.ensure_calls == [] and llm.calls == []
    assert "get_engine" not in env.names()  # nothing to transcribe, no model load
    assert result.warnings == ["В записи не найдено речи."]
    assert result.notes_md is not None
    assert "Конспект не составлен" in result.notes_md.read_text(encoding="utf-8")
    assert result.transcript_txt.read_text(encoding="utf-8") == "\n"


def test_pdf_failure_is_a_warning_and_docx_is_kept(env, monkeypatch):
    def broken_pdf(text, out_path, title):
        raise export.ExportError(f"Не удалось создать PDF «{out_path.name}» (нет движка).")

    monkeypatch.setattr(export, "export_notes_pdf", broken_pdf)
    monkeypatch.setattr(export, "export_transcript_pdf", broken_pdf)
    pipeline, events, _llm = make(env)
    pipeline.run(env.audio)
    result = result_of(events)
    assert result.transcript_pdf is None and result.notes_pdf is None
    assert result.transcript_docx.is_file() and result.notes_docx.is_file()
    assert result.warnings == ["Не удалось создать PDF «transcript.pdf» (нет движка).",
                               "Не удалось создать PDF «notes.pdf» (нет движка)."]


def test_unexpected_export_error_gets_file_name(env, monkeypatch):
    def broken(markdown, out_path, title):
        raise ValueError("bad xml")

    monkeypatch.setattr(export, "export_notes_docx", broken)
    pipeline, events, _llm = make(env)
    pipeline.run(env.audio)
    result = result_of(events)
    assert result.notes_docx is None and result.notes_pdf is not None
    assert result.warnings == ["Не удалось сохранить «notes.docx»: bad xml"]


def test_no_transcript_written_at_all_is_an_error(env, monkeypatch):
    def broken_text(path, text):
        raise RuntimeError(f"Не удалось сохранить «{path.name}»: диск полон")

    def broken_docx(text, out_path, title):
        raise export.ExportError("Не удалось сохранить «transcript.docx»: диск полон")

    monkeypatch.setattr(lp, "_write_text_atomic", broken_text)
    monkeypatch.setattr(export, "export_transcript_docx", broken_docx)
    pipeline, events, _llm = make(env)
    pipeline.run(env.audio)
    terminal = events.terminal()
    assert len(terminal) == 1 and terminal[0][0] == "error"
    assert terminal[0][1].startswith("Ошибка на этапе «Экспорт»: Не удалось сохранить расшифровку")


# --------------------------------------------------------------------------- errors & cancellation

def test_stage_exception_fires_only_on_error(env):
    env.preprocess_error = ValueError("плохие данные")
    pipeline, events, _llm = make(env)
    pipeline.run(env.audio)
    assert events.terminal() == [("error", "Ошибка на этапе «Предобработка»: плохие данные")]
    assert [s for (s,) in events.of("stage")] == ["preprocessing"]


def test_missing_file_fires_on_error(env, tmp_path):
    pipeline, events, _llm = make(env)
    pipeline.run(tmp_path / "raw_20260101_000000.wav")
    terminal = events.terminal()
    assert len(terminal) == 1 and terminal[0][0] == "error"
    assert "Файл не найден" in terminal[0][1]


def test_stt_load_failure_message(env, monkeypatch):
    def boom(checkpoint, device=None):
        raise ImportError("No module named 'gigaam'")

    monkeypatch.setattr(lp, "get_engine", boom)
    pipeline, events, _llm = make(env)
    pipeline.run(env.audio)
    assert events.terminal() == [(
        "error",
        "Ошибка на этапе «Транскрибация»: Не удалось загрузить модель распознавания речи v3_e2e_ctc: "
        "No module named 'gigaam'",
    )]


def test_memory_error_message(env):
    env.preprocess_error = MemoryError()
    pipeline, events, _llm = make(env)
    pipeline.run(env.audio)
    assert events.terminal() == [("error", "Ошибка на этапе «Предобработка»: недостаточно памяти")]


def test_cancel_during_diarization_fires_only_on_cancelled(env):
    pipeline, events, _llm = make(env)

    def cancel_now(on_progress):
        pipeline.cancel()
        on_progress(0.6)  # the pipeline's progress callback raises Cancelled
        raise AssertionError("diarization should have been interrupted")

    env.diarize_hook = cancel_now
    pipeline.run(env.audio)
    assert events.terminal() == [("cancelled",)]
    assert "vad" not in env.names()  # a cancel is not a diarization failure
    assert [s for (s,) in events.of("stage")] == ["preprocessing", "diarization"]
    assert not (env.out_dir / "segments.json").exists()


def test_cancel_inside_structuring(env):
    pipeline, events, llm = make(env)

    class CancellingLLM(FakeLLM):
        def stream_generate(self, model, prompt, system=None, on_token=None, options=None,
                            cancel_event=None, **_kw):
            pipeline.cancel()
            assert cancel_event is not None and cancel_event.is_set()
            return ""

    pipeline.llm = CancellingLLM()
    pipeline.run(env.audio)
    assert events.terminal() == [("cancelled",)]
    # The transcription cache survives a cancel, so the next run is quick.
    assert (env.out_dir / "segments.json").is_file()


def test_cancel_before_run(env):
    pipeline, events, _llm = make(env)
    pipeline.cancel()
    assert pipeline.cancelled
    pipeline.run(env.audio)
    assert events.items == [("cancelled",)]
    assert env.calls == []


def test_cancel_from_another_thread(env, monkeypatch):
    pipeline, events, _llm = make(env)
    started = threading.Event()

    def slow_preprocess(audio, sample_rate, do_denoise=False, do_dereverb=False, on_progress=None,
                        on_warning=None, cancel_event=None, device=None):
        started.set()
        assert cancel_event.wait(5.0)
        raise Cancelled()

    monkeypatch.setattr(preprocessing, "preprocess", slow_preprocess)
    worker = threading.Thread(target=pipeline.run, args=(env.audio,))
    worker.start()
    assert started.wait(5.0)
    pipeline.cancel()
    worker.join(5.0)
    assert not worker.is_alive()
    assert events.terminal() == [("cancelled",)]


def test_callback_exceptions_do_not_break_the_run(env):
    def broken(*_args):
        raise RuntimeError("UI bug")

    done = []
    pipeline = LecturePipeline(stt_tier="light", llm_tier="light", llm_client=FakeLLM(),
                               on_stage=broken, on_progress=broken, on_warning=broken,
                               on_done=done.append)
    pipeline.run(env.audio)
    assert len(done) == 1 and isinstance(done[0], LectureResult)


# --------------------------------------------------------------------------- misc

def test_hf_token_argument_then_env(monkeypatch):
    monkeypatch.setenv("HF_TOKEN", "hf_env")
    assert LecturePipeline(llm_client=FakeLLM()).hf_token == "hf_env"
    assert LecturePipeline(hf_token=" hf_arg ", llm_client=FakeLLM()).hf_token == "hf_arg"
    monkeypatch.delenv("HF_TOKEN")
    assert LecturePipeline(llm_client=FakeLLM()).hf_token == ""


def test_hf_token_is_passed_to_diarize(env):
    pipeline, events, _llm = make(env, hf_token="hf_secret")
    pipeline.run(env.audio)
    result_of(events)
    assert ("diarize", "hf_secret") in env.calls
    cache_text = (env.out_dir / "segments.json").read_text(encoding="utf-8")
    assert "hf_secret" not in cache_text


def test_recording_label(tmp_path):
    assert recording_label(tmp_path / "raw_20261003_201235.wav") == "03.10.2026 20:12"
    assert recording_label(tmp_path / "import_Лекция_5_матан_20261006_101500.wav") == "Лекция 5 матан"
    other = tmp_path / "lecture.wav"
    other.write_bytes(b"")
    import os
    os.utime(other, (1_700_000_000, 1_700_000_000))
    import datetime
    expected = datetime.datetime.fromtimestamp(1_700_000_000).strftime("%d.%m.%Y %H:%M")
    assert recording_label(other) == expected


def test_import_titles_and_title_hint(env, tmp_path, monkeypatch):
    from app.lecture import structuring

    hints = []
    real = structuring.structure_lecture

    def spy(*args, **kwargs):
        hints.append(kwargs.get("title_hint"))
        return real(*args, **kwargs)

    monkeypatch.setattr(structuring, "structure_lecture", spy)
    imported = tmp_path / "import_Матанализ_лекция_3_20261006_101500.wav"
    sf.write(str(imported), np.zeros(3 * SR, dtype=np.float32), SR, subtype="PCM_16")
    pipeline, events, llm = make(env)
    pipeline.run(imported)
    result_of(events)
    titles = {kind: title for kind, _name, title in env.pdf_calls}
    assert titles["notes"] == "Конспект лекции — Матанализ лекция 3"
    # How the hint is worded in the prompt is structuring's business.
    assert hints == ["Матанализ лекция 3"]
    assert "Матанализ лекция 3" in llm.calls[0]["prompt"]


def test_cpu_device_reaches_every_stage(env):
    pipeline, events, llm = make(env, device="cpu")
    pipeline.run(env.audio)
    result_of(events)
    assert env.devices == {"preprocess": "cpu", "diarize": "cpu", "stt": "cpu"}
    assert env.planned == [("lecture", "light", "cpu", None)]
    assert llm.calls[0]["options"]["num_gpu"] == 0
    assert env.released == []  # nothing on the GPU to give back


def test_gpu_device_defaults(env):
    pipeline, events, llm = make(env, llm_tier="ultra")
    pipeline.run(env.audio)
    result_of(events)
    # None = the torch stages pick CUDA themselves when it is available.
    assert env.devices == {"preprocess": None, "diarize": None, "stt": "gpu"}
    assert env.planned == [("lecture", "ultra", "gpu", None)]
    assert llm.calls[0]["model"] == llm_spec("lecture", "ultra").ollama_tag
    # Released right after transcription, and once more (a no-op) before structuring.
    assert env.released == [("v3_e2e_ctc", "gpu"), ("v3_e2e_ctc", "gpu")]


def test_options_planning_failure_falls_back_to_static(env, monkeypatch):
    def broken(*_a, **_kw):
        raise RuntimeError("nvidia-smi exploded")

    monkeypatch.setattr(lp, "plan_llm_options", broken)
    pipeline, events, llm = make(env)
    pipeline.run(env.audio)
    result_of(events)
    assert llm.calls[0]["options"] == llm_options("lecture", "light", "gpu")


def test_constants_match_spec():
    assert STAGES == ["preprocessing", "diarization", "transcription", "structuring", "export"]
    assert STAGE_LABELS == {"preprocessing": "Предобработка", "diarization": "Диаризация",
                            "transcription": "Транскрибация", "structuring": "Структурирование",
                            "export": "Экспорт"}
    assert CACHE_VERSION == 1


def test_default_llm_client_uses_settings_host(monkeypatch):
    class FakeSettings:
        def get(self, key, default=None):
            return {"ollama.host": "127.0.0.1:12345"}.get(key, default)

    monkeypatch.setattr(lp, "get_settings", lambda: FakeSettings())
    assert LecturePipeline().llm.host == "http://127.0.0.1:12345"


def test_fillers_only_skip_the_llm_and_vram_plan(env, monkeypatch):
    """Speech without lecture material: structure_lecture writes its stub, so the
    pipeline must not load (or plan VRAM for) a multi-GB model just to skip it."""
    from app.lecture import structuring

    monkeypatch.setattr(structuring, "MIN_CONTENT_WORDS", 1000)
    llm = FakeLLM()
    llm.ready_error = OllamaUnavailableError("не запущена")  # must not even be asked
    pipeline, events, _llm = make(env, llm=llm)
    pipeline.run(env.audio)
    result = result_of(events)
    assert llm.ensure_calls == [] and llm.calls == [] and env.planned == []
    assert result.notes_md is not None and result.warnings == []


# --------------------------------------------------------------------------- RAM guard (load limit)

class RamLLM(FakeLLM):
    """FakeLLM with what the RAM guard asks Ollama: sizes, loaded models, installed tags."""

    def __init__(self, installed):
        super().__init__()
        self.installed = set(installed)
        self.keep_alives = []

    def model_sizes(self):
        return {llm_spec("lecture", "power").ollama_tag: 5 * 1024 ** 3,
                llm_spec("lecture", "light").ollama_tag: int(2.5 * 1024 ** 3)}

    def running_models(self, timeout=2.0):
        return []

    def has_model(self, tag):
        return tag in self.installed

    def stream_generate(self, model, prompt, system=None, on_token=None, options=None, cancel_event=None,
                        keep_alive=None, **_kw):
        self.keep_alives.append(keep_alive)
        return super().stream_generate(model, prompt, system, on_token, options, cancel_event)


@pytest.fixture
def short_ram(monkeypatch):
    """16 GB of RAM with 7 GB in use: gemma (≈ 11 GB) doesn't fit under 95 %, T-lite (≈ 6.8 GB) does."""
    from app.core import resources

    monkeypatch.setattr(resources, "_installed", True)
    monkeypatch.setattr(resources, "_sizes", {})
    monkeypatch.setattr(resources, "RAM_POLL_S", 0.01)
    monkeypatch.setattr(resources, "memory_status", lambda: resources.MemoryStatus(16 * 1024 ** 3, 9 * 1024 ** 3))
    monkeypatch.setattr(lp, "RAM_WAIT_S", 0.0)
    return resources


def test_notes_fall_back_to_a_smaller_model_when_ram_is_short(env, short_ram):
    power = llm_spec("lecture", "power").ollama_tag
    llm = RamLLM(installed={power})
    pipeline, events, _llm = make(env, llm=llm, llm_tier="ultra")
    pipeline.run(env.audio)
    result = result_of(events)
    assert result.notes_docx is not None
    assert result.warnings == ["Не хватает оперативной памяти для модели «Очень мощная» — "
                               "конспект составлен моделью «Мощная»."]
    assert {call["model"] for call in llm.calls} == {power}
    assert llm.keep_alives[-1] == "0"  # the notes model leaves RAM right after the last chunk


def test_notes_are_skipped_when_no_model_fits(env, short_ram, monkeypatch):
    monkeypatch.setattr(short_ram, "memory_status",
                        lambda: short_ram.MemoryStatus(16 * 1024 ** 3, 2 * 1024 ** 3))
    llm = RamLLM(installed=set())
    pipeline, events, _llm = make(env, llm=llm, llm_tier="power")
    pipeline.run(env.audio)
    result = result_of(events)
    assert result.notes_docx is None and result.transcript_docx is not None
    assert result.warnings[0].startswith("Конспект не составлен. Не хватает оперативной памяти, чтобы составить "
                                         "конспект")
    assert llm.calls == []


def test_last_chunk_keeps_the_model_when_potok_uses_it(env, short_ram):
    light = llm_spec("lecture", "light").ollama_tag
    owner = object()
    short_ram.set_instant_model(light, owner)
    try:
        llm = RamLLM(installed={light})
        pipeline, events, _llm = make(env, llm=llm, llm_tier="light")
        pipeline.run(env.audio)
        result_of(events)
        assert llm.keep_alives[-1] is None  # the client default: «Поток» answers with it
    finally:
        short_ram.set_instant_model(None, owner)


def test_ram_switch_between_chunks_moves_the_hold_to_the_new_model(monkeypatch):
    from contextlib import ExitStack

    from app.core import resources

    monkeypatch.setattr(resources, "_held", {})
    monkeypatch.setattr(resources, "_listening", set())
    monkeypatch.setattr(resources, "ensure_ram_for_model", lambda *a, **kw: resources.RamCheck(False))
    pipeline = LecturePipeline(llm_tier="ultra", llm_client=FakeLLM())
    monkeypatch.setattr(pipeline, "_tier_for_ram", lambda result, below=None: "power")
    monkeypatch.setattr(pipeline, "_llm_options", lambda tier=None: {})
    ultra, power = llm_spec("lecture", "ultra").ollama_tag, llm_spec("lecture", "power").ollama_tag
    hold = ExitStack()
    hold.enter_context(resources.hold_model(ultra))
    current = {"tier": "ultra", "hold": hold}
    model, _options = pipeline._recheck_ram(1, current, result=None)
    assert model == power and current["tier"] == "power"
    # The model now in use can't be unloaded by another task's RAM check; the old one can.
    assert resources.protected_models() == {resources.canonical(power)}
    current["hold"].close()
    assert resources.protected_models() == set()
