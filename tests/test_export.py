"""Unit tests for app.lecture.export: markdown subset parsing, inline runs,
DOCX styling, escaped HTML and both PDF backends (no GPU/network needed).

The notes must look like plain, neat student notes: near-black text, bold
headings, no accent colour, no shaded boxes/borders/tables."""
import re
import threading
import zipfile
from pathlib import Path

import pytest
from docx import Document

from app.lecture import export
from app.lecture.export import Block, ExportError, inline_spans, parse_markdown

NOTES_MD = """## Краткое содержание

Лекция о **пределах** и *производных* <script>alert(1)</script> & co.

## Оглавление

1. Понятие предела [00:00–04:10]
2. Производная функции [04:10–09:30]

## 1. Понятие предела [00:00–04:10]

- **Предел** функции — число L.
- Односторонние пределы: *левый* и **правый *важный***.

> Уточнение по ходу лекции: предел может не существовать.

### Пример

Сравнение a < b & c > d.

## 2. Производная функции [04:10–09:30]

1. Определение через предел.
2. Геометрический смысл.

> Обычная цитата.
"""

TRANSCRIPT = "\n".join([
    "[00:00] Лектор: Сегодня мы поговорим о <b>пределах</b>.",
    "[00:12] Аудитория: А это будет на экзамене?",
    "",
    "Строка без формата",
    "[1:02:03] Лектор: Через час & позже.",
])


@pytest.fixture(scope="module")
def qt_app():
    """One Qt application for the process, created on the main thread.

    A QApplication (not just QGuiApplication) so other test modules that need
    widgets can reuse it."""
    from PyQt6.QtCore import QCoreApplication
    from PyQt6.QtWidgets import QApplication

    app = QCoreApplication.instance()
    if app is None:
        app = QApplication([])
    yield app


@pytest.fixture
def force_qt_backend(monkeypatch):
    monkeypatch.setattr(export, "_import_weasyprint", lambda: None)
    export.pdf_backend.cache_clear()
    yield
    export.pdf_backend.cache_clear()


def _run_in_thread(fn, *args):
    result = {}

    def target():
        try:
            fn(*args)
        except BaseException as exc:  # noqa: BLE001 — re-raised in the test thread
            result["error"] = exc

    t = threading.Thread(target=target)
    t.start()
    t.join(60)
    assert not t.is_alive(), "export thread hung"
    if "error" in result:
        raise result["error"]


def _pdf_text(path: Path) -> tuple:
    from PyQt6.QtPdf import QPdfDocument

    doc = QPdfDocument(None)
    doc.load(str(path))
    pages = doc.pageCount()
    text = "".join(doc.getAllText(i).text() for i in range(pages))
    doc.close()
    return pages, text


# --------------------------------------------------------------------------
# parse_markdown
# --------------------------------------------------------------------------

def test_parse_headings_lists_quotes_and_callouts():
    blocks = parse_markdown(NOTES_MD)
    kinds = [b.kind for b in blocks]
    assert kinds == ["h2", "p", "h2", "ol", "h2", "ul", "callout", "h3", "p", "h2", "ol", "quote"]
    assert blocks[0].text == "Краткое содержание"
    assert blocks[3].items == ["Понятие предела [00:00–04:10]", "Производная функции [04:10–09:30]"]
    assert blocks[5].items[0] == "**Предел** функции — число L."
    assert blocks[6].text == "Уточнение по ходу лекции: предел может не существовать."
    assert blocks[11].text == "Обычная цитата."


def test_parse_heading_levels_and_trailing_hashes():
    blocks = parse_markdown("# A\n## B ##\n### C\n#### D\n####### not a heading\n## C# basics")
    assert [(b.kind, b.text) for b in blocks[:4]] == [("h1", "A"), ("h2", "B"), ("h3", "C"), ("h3", "D")]
    assert blocks[4].kind == "p" and blocks[4].text == "####### not a heading"
    assert (blocks[5].kind, blocks[5].text) == ("h2", "C# basics")


def test_parse_paragraph_joining_and_separation():
    blocks = parse_markdown("первая строка\nвторая строка\n\nновый абзац")
    assert [(b.kind, b.text) for b in blocks] == [("p", "первая строка вторая строка"), ("p", "новый абзац")]


def test_parse_lists_loose_star_bullets_and_switch():
    md = "- a\n\n- b\n* c\n+ d\n1. one\n2) two\n- e"
    blocks = parse_markdown(md)
    assert [(b.kind, b.items) for b in blocks] == [
        ("ul", ["a", "b", "c", "d"]), ("ol", ["one", "two"]), ("ul", ["e"])
    ]


def test_parse_bold_line_is_not_a_bullet_and_hr_fences_ignored():
    md = "**Жирная** строка\n\n* * *\n\n```python\nкод\n```\n---\n*курсив* текст"
    blocks = parse_markdown(md)
    assert [(b.kind, b.text) for b in blocks] == [
        ("p", "**Жирная** строка"), ("p", "код"), ("p", "*курсив* текст")
    ]


def test_parse_list_item_continuation_and_quote_lazy_continuation():
    blocks = parse_markdown("- пункт\n  продолжение\n> цитата\nпродолжение цитаты\n\nабзац")
    assert blocks[0].items == ["пункт продолжение"]
    assert (blocks[1].kind, blocks[1].text) == ("quote", "цитата\nпродолжение цитаты")
    assert (blocks[2].kind, blocks[2].text) == ("p", "абзац")


def test_parse_callout_variants_split_into_own_blocks():
    md = ("> **Уточнение по ходу лекции:** первое\n"
          "> Уточнение по ходу лекции — второе\n"
          ">\n"
          "> просто цитата\n\n"
          "Уточнение по ходу лекции: абзац-врезка")
    blocks = parse_markdown(md)
    assert [b.kind for b in blocks] == ["callout", "callout", "quote", "callout"]
    assert export.split_callout(blocks[0].text) == ("Уточнение по ходу лекции", "первое")
    assert export.split_callout(blocks[1].text) == ("Уточнение по ходу лекции", "второе")
    assert export.split_callout("обычный текст") == ("", "обычный текст")


def test_parse_empty_input():
    assert parse_markdown("") == []
    assert parse_markdown("\n\n  \n") == []
    assert parse_markdown(None) == []


def test_block_defaults():
    assert Block("p").items == [] and Block("p").text == ""


# --------------------------------------------------------------------------
# inline_spans
# --------------------------------------------------------------------------

@pytest.mark.parametrize("text, expected", [
    ("", []),
    ("plain", [("plain", False, False)]),
    ("a **b** c", [("a ", False, False), ("b", True, False), (" c", False, False)]),
    ("a *b* c", [("a ", False, False), ("b", False, True), (" c", False, False)]),
    ("***оба***", [("оба", True, True)]),
    ("**жирный *и курсив* текст**",
     [("жирный ", True, False), ("и курсив", True, True), (" текст", True, False)]),
    ("*курсив **и жирный** текст*",
     [("курсив ", False, True), ("и жирный", True, True), (" текст", False, True)]),
    ("2 * 3 * 4", [("2 * 3 * 4", False, False)]),
    ("**незакрытый", [("**незакрытый", False, False)]),
    ("a ** b", [("a ** b", False, False)]),
    (r"звёздочка \*не курсив\*", [("звёздочка *не курсив*", False, False)]),
    ("x < y & **a > b**", [("x < y & ", False, False), ("a > b", True, False)]),
])
def test_inline_spans(text, expected):
    assert inline_spans(text) == expected


def test_inline_spans_roundtrip_text_is_preserved():
    text = "Термин **предел** и *эпсилон-дельта* <tag> & «кавычки»"
    assert "".join(t for t, _b, _i in inline_spans(text)) == "Термин предел и эпсилон-дельта <tag> & «кавычки»"


# --------------------------------------------------------------------------
# DOCX
# --------------------------------------------------------------------------

def test_notes_docx_styles_and_runs(tmp_path):
    out = tmp_path / "notes.docx"
    export.export_notes_docx(NOTES_MD, out, "Конспект лекции — 05.10.2026 14:30")
    doc = Document(str(out))
    paras = [p for p in doc.paragraphs if p.text.strip()]
    styles = [p.style.name for p in paras]
    assert styles[0] == "Title" and paras[0].text == "Конспект лекции — 05.10.2026 14:30"
    assert styles.count("Heading 2") == 4
    assert "Heading 3" in styles
    assert styles.count("List Bullet") == 2
    assert styles.count("List Number") == 4
    assert "Intense Quote" not in styles
    assert "Quote" in styles
    full = "\n".join(p.text for p in paras)
    assert "**" not in full
    assert not any(p.text.startswith(("- ", "> ", "#")) for p in paras)
    # "1. " survives only as part of a numbered section heading, never as a list marker.
    assert not any(p.text.startswith("1. ") for p in paras if not p.style.name.startswith("Heading"))
    assert "<script>alert(1)</script> & co." in full  # literal text, not markup

    bullet = next(p for p in paras if p.text.startswith("Предел"))
    assert [(r.text, bool(r.bold)) for r in bullet.runs][:2] == [("Предел", True), (" функции — число L.", False)]
    nested = next(p for p in paras if p.text.startswith("Односторонние"))
    runs = [(r.text, bool(r.bold), bool(r.italic)) for r in nested.runs]
    assert ("левый", False, True) in runs
    assert ("важный", True, True) in runs

    # A legacy callout is an ordinary paragraph with a bold label: no box, no colour.
    callout = next(p for p in paras if p.text.startswith("Уточнение по ходу лекции"))
    assert callout.style.name == "Normal" and callout.runs[0].bold
    assert callout.text == "Уточнение по ходу лекции: предел может не существовать."
    assert all(r.font.color.rgb is None for r in callout.runs)


def test_notes_docx_numbered_lists_restart(tmp_path):
    out = tmp_path / "notes.docx"
    export.export_notes_docx("1. a\n2. b\n\nтекст\n\n1. c\n2. d", out, "T")
    doc = Document(str(out))
    numbered = [p for p in doc.paragraphs if p.style.name == "List Number"]
    num_ids = [p._p.pPr.numPr.numId.val for p in numbered]
    assert num_ids[0] == num_ids[1] and num_ids[2] == num_ids[3] and num_ids[0] != num_ids[2]
    numbering = doc.part.numbering_part.element
    for num_id in (num_ids[0], num_ids[2]):
        num = numbering.num_having_numId(num_id)
        assert num.lvlOverride_lst and num.lvlOverride_lst[0].startOverride.val == 1


_ACCENTS = ("0F6CBD", "EEF5FC", "4F81BD", "365F91", "17365D")
_USED_STYLES = ("Normal", "Title", "Heading 1", "Heading 2", "Heading 3", "List Bullet", "List Number", "Quote")


def _docx_xml(path, name):
    with zipfile.ZipFile(path) as z:
        return z.read(name).decode("utf-8")


@pytest.mark.parametrize("kind", ["notes", "transcript"])
def test_docx_is_plain_dark_without_accent_colour(tmp_path, kind):
    out = tmp_path / f"{kind}.docx"
    if kind == "notes":
        export.export_notes_docx("# Тема\n\n" + NOTES_MD, out, "Конспект лекции — 05.10.2026 14:30")
    else:
        export.export_transcript_docx(TRANSCRIPT, out, "Расшифровка лекции — 05.10.2026 14:30")
    doc = Document(str(out))
    from docx.oxml.ns import qn

    document_xml = _docx_xml(out, "word/document.xml")
    for accent in _ACCENTS:
        assert accent not in document_xml, accent
    assert "w:shd" not in document_xml and "w:pBdr" not in document_xml and "w:tbl" not in document_xml
    assert "0F6CBD" not in _docx_xml(out, "word/styles.xml")
    for name in _USED_STYLES:
        style = doc.styles[name]
        xml = style.element.xml
        assert "themeColor" not in xml and "w:pBdr" not in xml and "w:shd" not in xml, name
        assert not any(a in xml for a in _ACCENTS), name
        rfonts = style.element.rPr.rFonts if style.element.rPr is not None else None
        if rfonts is not None:
            assert rfonts.get(qn("w:asciiTheme")) is None and rfonts.get(qn("w:ascii")) == export.DOCX_FONT
    for name in ("Title", "Heading 1", "Heading 2", "Heading 3"):
        font = doc.styles[name].font
        assert font.bold and str(font.color.rgb) == "1B1B1B", name
    assert doc.styles["Normal"].font.size.pt == export.BODY_PT == 11
    # Every explicit run colour is a dark or neutral gray.
    for value in re.findall(r'<w:color w:val="([0-9A-Fa-f]{6})"', document_xml):
        r, g, b = (int(value[i:i + 2], 16) for i in (0, 2, 4))
        assert max(r, g, b) - min(r, g, b) <= 8, value
    assert doc.core_properties.author == "Нууус"


def test_notes_docx_empty_and_control_chars(tmp_path):
    out = tmp_path / "empty.docx"
    export.export_notes_docx("", out, "Пусто")
    assert "Конспект пуст." in [p.text for p in Document(str(out)).paragraphs]
    out2 = tmp_path / "ctrl.docx"
    export.export_notes_docx("## Заголовок\x01\n\n- пункт\x0b с мусором\x00", out2, "T\x02")
    texts = [p.text for p in Document(str(out2)).paragraphs]
    assert "Заголовок" in texts and "пункт с мусором" in texts


def test_transcript_docx_runs(tmp_path):
    out = tmp_path / "transcript.docx"
    export.export_transcript_docx(TRANSCRIPT, out, "Расшифровка лекции")
    doc = Document(str(out))
    paras = [p for p in doc.paragraphs if p.text.strip()]
    assert paras[0].style.name == "Title"
    lecturer = paras[1]
    tc, speaker, text = lecturer.runs
    assert tc.text.strip() == "00:00"
    assert str(tc.font.color.rgb) == "808080" and tc.font.size.pt == 8.5
    assert speaker.text == "Лектор: " and speaker.bold
    assert text.text == "Сегодня мы поговорим о <b>пределах</b>." and not text.italic
    assert lecturer.paragraph_format.left_indent is None

    audience = paras[2]
    assert audience.runs[1].text == "Аудитория: " and audience.runs[1].bold
    assert audience.runs[2].italic
    assert audience.paragraph_format.left_indent is not None and audience.paragraph_format.left_indent > 0

    plain = paras[3]
    assert plain.text == "Строка без формата" and len(plain.runs) == 1
    assert paras[4].runs[0].text.strip() == "1:02:03"


def test_transcript_docx_empty(tmp_path):
    out = tmp_path / "t.docx"
    export.export_transcript_docx("", out, "T")
    assert "Речь в записи не распознана." in [p.text for p in Document(str(out)).paragraphs]


def test_docx_locked_target_gives_russian_error(tmp_path, monkeypatch):
    out = tmp_path / "notes.docx"

    def locked(src, dst):
        raise PermissionError(13, "denied")

    monkeypatch.setattr(export.os, "replace", locked)
    with pytest.raises(ExportError, match="открыт в другой программе"):
        export.export_notes_docx("текст", out, "T")
    assert list(tmp_path.iterdir()) == []  # temp file cleaned up


# --------------------------------------------------------------------------
# HTML
# --------------------------------------------------------------------------

@pytest.mark.parametrize("qt", [False, True])
def test_notes_html_escapes_and_structure(qt):
    html = export.render_notes_html(NOTES_MD, "Заголовок <i>&</i>", qt=qt)
    assert "<script>" not in html and "&lt;script&gt;alert(1)&lt;/script&gt; &amp; co." in html
    assert "Заголовок &lt;i&gt;&amp;&lt;/i&gt;" in html
    assert "<strong>Предел</strong>" in html and "<em>левый</em>" in html
    assert "<strong>правый </strong><strong><em>важный</em></strong>" in html
    assert "a &lt; b &amp; c &gt; d" in html
    assert html.count("<ol>") == 2 and html.count("<ul>") == 1
    assert '"Calibri", "Segoe UI", "DejaVu Sans", Arial, sans-serif' in html  # not entity-escaped inside <style>
    assert ("@page" in html) is not qt
    assert "<p><strong>Уточнение по ходу лекции:</strong> предел может не существовать.</p>" in html


@pytest.mark.parametrize("qt", [False, True])
@pytest.mark.parametrize("kind", ["notes", "transcript"])
def test_html_is_plain_without_accent_colour_or_boxes(qt, kind):
    if kind == "notes":
        html = export.render_notes_html(NOTES_MD, "Конспект", qt=qt)
    else:
        html = export.render_transcript_html(TRANSCRIPT, "Расшифровка", qt=qt)
    lowered = html.lower()
    for bad in ("0f6cbd", "eef5fc", "bgcolor", "<table", "border", "background", 'class="callout'):
        assert bad not in lowered, bad
    for value in re.findall(r"#([0-9a-f]{6})\b", lowered):
        r, g, b = (int(value[i:i + 2], 16) for i in (0, 2, 4))
        assert max(r, g, b) - min(r, g, b) <= 8, value  # only grays
    assert "font-size: 11pt" in html


def test_transcript_html_escapes_and_marks_audience():
    html = export.render_transcript_html(TRANSCRIPT, "T")
    assert "&lt;b&gt;пределах&lt;/b&gt;" in html and "<b>" not in html
    assert 'class="line aud"' in html
    assert '<span class="tc">00:12</span>' in html
    assert "Через час &amp; позже." in html


# --------------------------------------------------------------------------
# PDF
# --------------------------------------------------------------------------

def test_pdf_backend_probe_is_cached():
    export.pdf_backend.cache_clear()
    first = export.pdf_backend()
    assert first in ("weasyprint", "qt", "none")
    assert export.pdf_backend() == first
    assert export.pdf_backend.cache_info().hits >= 1


def test_pdf_weasyprint_backend(tmp_path, qt_app):
    export.pdf_backend.cache_clear()
    if export.pdf_backend() != "weasyprint":
        pytest.skip("WeasyPrint is not available here")
    out = tmp_path / "notes.pdf"
    export.export_notes_pdf(NOTES_MD, out, "Конспект лекции")
    assert out.read_bytes()[:5] == b"%PDF-"
    pages, text = _pdf_text(out)
    assert pages >= 1 and "Конспект лекции" in text and "Уточнение по ходу лекции" in text
    assert "**" not in text
    out2 = tmp_path / "transcript.pdf"
    export.export_transcript_pdf(TRANSCRIPT, out2, "Расшифровка")
    assert out2.read_bytes()[:5] == b"%PDF-"


def test_pdf_qt_backend_from_worker_thread(tmp_path, qt_app, force_qt_backend):
    assert export.pdf_backend() == "qt"
    notes = tmp_path / "notes.pdf"
    transcript = tmp_path / "transcript.pdf"
    _run_in_thread(export.export_notes_pdf, NOTES_MD * 3, notes, "Конспект лекции")
    _run_in_thread(export.export_transcript_pdf, "\n".join([TRANSCRIPT] * 40), transcript, "Расшифровка")
    for path in (notes, transcript):
        assert path.read_bytes()[:5] == b"%PDF-"
        assert not (tmp_path / f"{path.stem}.tmp.pdf").exists()
    pages, text = _pdf_text(notes)
    assert pages >= 1
    assert "Конспект лекции" in text and "Понятие предела" in text and "**" not in text
    pages, text = _pdf_text(transcript)
    assert pages >= 2 and "Аудитория:" in text and "1 / " in text  # paginated with footer numbers

    from PyQt6.QtPdf import QPdfDocument

    doc = QPdfDocument(None)
    doc.load(str(notes))
    try:
        assert doc.metaData(QPdfDocument.MetaDataField.Creator) == export.APP_AUTHOR == "Нууус"
    finally:
        doc.close()


def test_pdf_weasyprint_runtime_failure_falls_back_to_qt(tmp_path, qt_app, monkeypatch):
    monkeypatch.setattr(export, "pdf_backend", lambda: "weasyprint")

    def broken(html, out_path):
        raise RuntimeError("cairo exploded")

    monkeypatch.setattr(export, "_pdf_weasyprint", broken)
    out = tmp_path / "notes.pdf"
    export.export_notes_pdf(NOTES_MD, out, "T")
    assert out.read_bytes()[:5] == b"%PDF-"


def test_pdf_no_backend_raises_export_error(tmp_path, monkeypatch):
    monkeypatch.setattr(export, "pdf_backend", lambda: "none")
    with pytest.raises(ExportError, match="Не удалось создать PDF"):
        export.export_notes_pdf("текст", tmp_path / "x.pdf", "T")
    assert not (tmp_path / "x.pdf").exists()


def test_pdf_qt_without_app_in_worker_thread_raises(tmp_path, monkeypatch, force_qt_backend):
    from PyQt6.QtCore import QCoreApplication

    try:
        monkeypatch.setattr(QCoreApplication, "instance", staticmethod(lambda: None))
    except (TypeError, AttributeError):
        pytest.skip("cannot patch QCoreApplication.instance in this PyQt build")
    with pytest.raises(ExportError, match="запущенное приложение"):
        _run_in_thread(export.export_notes_pdf, "текст", tmp_path / "x.pdf", "T")
